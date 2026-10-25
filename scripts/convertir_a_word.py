#!/usr/bin/env python3
"""
Script de Conversión de Manuales Markdown a Microsoft Word (.docx)
Proyecto: ConectaTec · Triaje Inteligente de Incidentes ISP (Caso N.º 03 · Grupo 3 UNIANDES)
Seminario de Titulación: Ingeniería de Software con IA Generativa - 2026

Convierte todos los archivos .md de la carpeta material_estudio/ a documentos Word (.docx)
con diseño editorial profesional (paleta institucional, tipografía jerárquica, tablas con estilos,
bloques de código sombreados y cajas de notas/citas).
"""

import os
import sys
import re
import glob

# Forzar codificación UTF-8 en salida de consola Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
CARPETA_MATERIAL = os.path.join(RAIZ, "material_estudio")

# Colores institucionales
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A (Títulos principales)
COLOR_AZUL = RGBColor(37, 99, 235)       # #2563EB (Subtítulos H2)
COLOR_GRIS_OSCURO = RGBColor(55, 65, 81) # #374151 (H3 y texto)
COLOR_TEXTO = RGBColor(31, 41, 55)       # #1F2937 (Cuerpo)
COLOR_MUTED = RGBColor(107, 114, 128)    # #6B7280 (Metadatos)
COLOR_CODIGO = RGBColor(220, 38, 38)     # #DC2626 (Código inline rojo suave)

HEX_ENCABEZADO_TABLA = "1E3A8A"
HEX_FONDO_TABLA_ALT = "F8FAFC"
HEX_BORDE_TABLA = "CBD5E1"
HEX_FONDO_CODIGO = "F1F5F9"
HEX_BORDE_CODIGO = "CBD5E1"
HEX_FONDO_CITA = "EFF6FF"
HEX_BORDE_CITA = "3B82F6"


def aplicar_sombreado_celda(celda, color_hex):
    """Aplica color de fondo hexadecimal a una celda de tabla."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    celda._tc.get_or_add_tcPr().append(shading_elm)


def aplicar_margenes_celda(celda, top=120, bottom=120, left=180, right=180):
    """Configura márgenes internos (padding) de una celda en dxa."""
    tcPr = celda._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)


def configurar_bordes_tabla(tabla, color=HEX_BORDE_TABLA):
    """Establece bordes sutiles y limpios a una tabla."""
    tblPr = tabla._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def agregar_texto_con_formato(parrafo, texto, fuente="Calibri", tamano_pt=11, color_rgb=COLOR_TEXTO, es_negrita_base=False, es_cursiva_base=False):
    """
    Parsea tokens inline de Markdown (**negrita**, *cursiva*, `código`) y los inserta como runs.
    """
    if not texto:
        return

    # Expresión regular para separar tokens Markdown inline
    patron = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    partes = patron.split(texto)

    for parte in partes:
        if not parte:
            continue
        run = parrafo.add_run()
        run.font.name = fuente
        run.font.size = Pt(tamano_pt)
        run.font.color.rgb = color_rgb
        run.bold = es_negrita_base
        run.italic = es_cursiva_base

        if parte.startswith("**") and parte.endswith("**") and len(parte) >= 4:
            run.text = parte[2:-2]
            run.bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) >= 2:
            run.text = parte[1:-1]
            run.italic = True
        elif parte.startswith("`") and parte.endswith("`") and len(parte) >= 2:
            run.text = parte[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(tamano_pt * 0.92)
            run.font.color.rgb = COLOR_CODIGO
        else:
            run.text = parte


def crear_documento_word(ruta_md, ruta_docx):
    """Convierte un archivo Markdown estructurado a un documento .docx profesional."""
    doc = docx.Document()

    # Configurar márgenes de página (1 pulgada = 2.54 cm)
    for seccion in doc.sections:
        seccion.top_margin = Inches(1.0)
        seccion.bottom_margin = Inches(1.0)
        seccion.left_margin = Inches(1.0)
        seccion.right_margin = Inches(1.0)

        # Encabezado institucional
        encabezado = seccion.header
        p_enc = encabezado.paragraphs[0]
        p_enc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_enc = p_enc.add_run("ConectaTec · Grupo 3 UNIANDES | Material de Estudio y Preparación para Tribunal")
        r_enc.font.name = "Calibri"
        r_enc.font.size = Pt(8.5)
        r_enc.font.color.rgb = COLOR_MUTED

        # Pie de página institucional
        pie = seccion.footer
        p_pie = pie.paragraphs[0]
        p_pie.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r_pie = p_pie.add_run("Caso N.º 03 · Triaje Inteligente ISP — Documento Interno de Estudio (Confidencial)")
        r_pie.font.name = "Calibri"
        r_pie.font.size = Pt(8.5)
        r_pie.font.color.rgb = COLOR_MUTED

    with open(ruta_md, "r", encoding="utf-8") as f:
        lineas = f.readlines()

    en_bloque_codigo = False
    lineas_codigo = []
    lenguaje_codigo = ""

    en_tabla = False
    filas_tabla = []

    def vaciar_tabla():
        nonlocal en_tabla, filas_tabla
        if not filas_tabla:
            en_tabla = False
            return

        num_cols = max(len(fila) for fila in filas_tabla)
        if num_cols == 0:
            en_tabla = False
            filas_tabla = []
            return

        tabla = doc.add_table(rows=len(filas_tabla), cols=num_cols)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        configurar_bordes_tabla(tabla)

        for i_fila, fila_datos in enumerate(filas_tabla):
            fila_doc = tabla.rows[i_fila]
            es_header = (i_fila == 0)

            for i_col in range(num_cols):
                celda = fila_doc.cells[i_col]
                texto_celda = fila_datos[i_col].strip() if i_col < len(fila_datos) else ""

                aplicar_margenes_celda(celda, top=140, bottom=140, left=160, right=160)

                if es_header:
                    aplicar_sombreado_celda(celda, HEX_ENCABEZADO_TABLA)
                    p_celda = celda.paragraphs[0]
                    p_celda.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_celda.paragraph_format.space_before = Pt(2)
                    p_celda.paragraph_format.space_after = Pt(2)
                    run = p_celda.add_run(texto_celda)
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    if i_fila % 2 == 1:
                        aplicar_sombreado_celda(celda, HEX_FONDO_TABLA_ALT)
                    p_celda = celda.paragraphs[0]
                    p_celda.paragraph_format.space_before = Pt(2)
                    p_celda.paragraph_format.space_after = Pt(2)
                    agregar_texto_con_formato(p_celda, texto_celda, fuente="Calibri", tamano_pt=9.5, color_rgb=COLOR_TEXTO)

        # Espacio después de la tabla
        p_post = doc.add_paragraph()
        p_post.paragraph_format.space_before = Pt(0)
        p_post.paragraph_format.space_after = Pt(6)

        filas_tabla = []
        en_tabla = False

    def vaciar_codigo():
        nonlocal en_bloque_codigo, lineas_codigo, lenguaje_codigo
        if not lineas_codigo:
            en_bloque_codigo = False
            return

        # Renderizar bloque de código como una celda de tabla estilizada
        tabla_cod = doc.add_table(rows=1, cols=1)
        tabla_cod.alignment = WD_TABLE_ALIGNMENT.CENTER
        celda = tabla_cod.rows[0].cells[0]
        aplicar_sombreado_celda(celda, HEX_FONDO_CODIGO)
        aplicar_margenes_celda(celda, top=140, bottom=140, left=180, right=180)

        # Borde sutil alrededor del bloque de código
        tblPr = tabla_cod._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDE_CODIGO}"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDE_CODIGO}"/>
                <w:left w:val="single" w:sz="12" w:space="0" w:color="{HEX_BORDE_CITA}"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDE_CODIGO}"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

        for i, linea_c in enumerate(lineas_codigo):
            if i == 0:
                p_c = celda.paragraphs[0]
            else:
                p_c = celda.add_paragraph()
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(1)
            p_c.paragraph_format.line_spacing = 1.05

            run = p_c.add_run(linea_c)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = COLOR_TEXTO

        p_post = doc.add_paragraph()
        p_post.paragraph_format.space_before = Pt(0)
        p_post.paragraph_format.space_after = Pt(6)

        lineas_codigo = []
        en_bloque_codigo = False

    i = 0
    while i < len(lineas):
        linea_raw = lineas[i]
        linea = linea_raw.rstrip("\r\n")

        # 1. Detección de bloques de código
        if linea.strip().startswith("```"):
            if en_bloque_codigo:
                # Cierre del bloque de código
                vaciar_codigo()
            else:
                if en_tabla:
                    vaciar_tabla()
                en_bloque_codigo = True
                lenguaje_codigo = linea.strip()[3:].strip()
                lineas_codigo = []
            i += 1
            continue

        if en_bloque_codigo:
            lineas_codigo.append(linea)
            i += 1
            continue

        # 2. Detección de tablas Markdown (| col | col |)
        if linea.strip().startswith("|") and "|" in linea.strip()[1:]:
            # Comprobar si es línea separadora (|---|---|)
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', linea):
                # Es el divisor de la cabecera, lo ignoramos para los datos
                i += 1
                continue
            en_tabla = True
            # Parsear celdas
            partes_fila = [c.strip() for c in linea.strip().strip("|").split("|")]
            filas_tabla.append(partes_fila)
            i += 1
            continue
        elif en_tabla:
            # Termina la tabla
            vaciar_tabla()

        # 3. Línea horizontal (--- o ***)
        if re.match(r'^\s*([-_*]){3,}\s*$', linea):
            p_hr = doc.add_paragraph()
            p_hr.paragraph_format.space_before = Pt(6)
            p_hr.paragraph_format.space_after = Pt(6)
            run_hr = p_hr.add_run("_________________________________________________________________________________")
            run_hr.font.name = "Calibri"
            run_hr.font.size = Pt(8)
            run_hr.font.color.rgb = COLOR_MUTED
            i += 1
            continue

        # 4. Líneas vacías
        if not linea.strip():
            i += 1
            continue

        # 5. Encabezados (#, ##, ###, ####)
        if linea.startswith("# "):
            titulo_texto = linea[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            agregar_texto_con_formato(p, titulo_texto, fuente="Calibri", tamano_pt=20, color_rgb=COLOR_NAVY, es_negrita_base=True)
            i += 1
            continue

        if linea.startswith("## "):
            h2_texto = linea[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            agregar_texto_con_formato(p, h2_texto, fuente="Calibri", tamano_pt=14, color_rgb=COLOR_AZUL, es_negrita_base=True)
            i += 1
            continue

        if linea.startswith("### "):
            h3_texto = linea[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            agregar_texto_con_formato(p, h3_texto, fuente="Calibri", tamano_pt=12, color_rgb=COLOR_GRIS_OSCURO, es_negrita_base=True)
            i += 1
            continue

        if linea.startswith("#### "):
            h4_texto = linea[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            agregar_texto_con_formato(p, h4_texto, fuente="Calibri", tamano_pt=11, color_rgb=COLOR_GRIS_OSCURO, es_negrita_base=True)
            i += 1
            continue

        # 6. Citas / Callouts (> texto)
        if linea.strip().startswith(">"):
            contenido_cita = linea.strip().lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run_barra = p.add_run("┃ ")
            run_barra.font.name = "Calibri"
            run_barra.font.bold = True
            run_barra.font.color.rgb = COLOR_AZUL
            agregar_texto_con_formato(p, contenido_cita, fuente="Calibri", tamano_pt=10.5, color_rgb=COLOR_GRIS_OSCURO, es_cursiva_base=True)
            i += 1
            continue

        # 7. Listas desordenadas (* o -)
        m_bullet = re.match(r'^(\s*)[*-]\s+(.*)', linea)
        if m_bullet:
            espacios, contenido = m_bullet.groups()
            nivel_indent = len(espacios) // 2
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25 * (nivel_indent + 1))
            p.paragraph_format.first_line_indent = Inches(-0.15)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            run_b = p.add_run("•  ")
            run_b.font.name = "Calibri"
            run_b.font.size = Pt(10)
            run_b.font.color.rgb = COLOR_AZUL

            agregar_texto_con_formato(p, contenido, fuente="Calibri", tamano_pt=10.5, color_rgb=COLOR_TEXTO)
            i += 1
            continue

        # 8. Listas ordenadas (1. , 2. )
        m_num = re.match(r'^(\s*)(\d+)\.\s+(.*)', linea)
        if m_num:
            espacios, numero, contenido = m_num.groups()
            nivel_indent = len(espacios) // 2
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25 * (nivel_indent + 1))
            p.paragraph_format.first_line_indent = Inches(-0.15)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            run_n = p.add_run(f"{numero}. ")
            run_n.font.name = "Calibri"
            run_n.font.size = Pt(10)
            run_n.font.bold = True
            run_n.font.color.rgb = COLOR_NAVY

            agregar_texto_con_formato(p, contenido, fuente="Calibri", tamano_pt=10.5, color_rgb=COLOR_TEXTO)
            i += 1
            continue

        # 9. Párrafo estándar
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        agregar_texto_con_formato(p, linea, fuente="Calibri", tamano_pt=11, color_rgb=COLOR_TEXTO)
        i += 1

    # Vaciar posibles pendientes
    if en_bloque_codigo:
        vaciar_codigo()
    if en_tabla:
        vaciar_tabla()

    doc.save(ruta_docx)
    # Crear también versión con extensión .doc solicitada
    ruta_doc = os.path.splitext(ruta_docx)[0] + ".doc"
    try:
        import shutil
        shutil.copyfile(ruta_docx, ruta_doc)
    except Exception:
        pass
    print(f"   ✅ Generado Word: {os.path.basename(ruta_docx)} y {os.path.basename(ruta_doc)}")


def main():
    print("=" * 70)
    print("🚀 CONVERTIDOR DE MANUALES MARKDOWN A MICROSOFT WORD (.DOCX)")
    print("   Proyecto: ConectaTec · Triaje Inteligente ISP (Grupo 3 UNIANDES)")
    print(f"   Carpeta Origen: {CARPETA_MATERIAL}")
    print("=" * 70)

    if not os.path.isdir(CARPETA_MATERIAL):
        print(f"❌ Error: La carpeta {CARPETA_MATERIAL} no existe.")
        sys.exit(1)

    patron_md = os.path.join(CARPETA_MATERIAL, "*.md")
    archivos_md = sorted(glob.glob(patron_md))

    if not archivos_md:
        print("ℹ️ No se encontraron archivos .md en la carpeta.")
        return

    print(f"📂 Se encontraron {len(archivos_md)} manuales en formato Markdown (.md).\n")

    convertidos = 0
    errores = 0

    for ruta_md in archivos_md:
        nombre_base = os.path.splitext(os.path.basename(ruta_md))[0]
        ruta_docx = os.path.join(CARPETA_MATERIAL, f"{nombre_base}.docx")
        
        print(f"⚙️ Procesando: {os.path.basename(ruta_md)} ...")
        try:
            crear_documento_word(ruta_md, ruta_docx)
            convertidos += 1
        except Exception as e:
            print(f"   ❌ Error al convertir {os.path.basename(ruta_md)}: {e}")
            errores += 1

    print("\n" + "=" * 70)
    print(f"🎉 Proceso finalizado:")
    print(f"   • Documentos Word creados exitosamente: {convertidos}/{len(archivos_md)}")
    if errores > 0:
        print(f"   • Errores detectados: {errores}")
    print("=" * 70)


if __name__ == "__main__":
    main()
